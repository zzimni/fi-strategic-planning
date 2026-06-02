"""
reporter.py
xlsx 리포트 생성 모듈
"""

import os
from datetime import datetime, date
from typing import Optional

import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter
from openpyxl.styles.numbers import FORMAT_NUMBER_COMMA_SEPARATED1

from filters import CODE_NAMES, parse_date

# ─────────────────────────────────────────────
# 색상 팔레트
# ─────────────────────────────────────────────

CLR_S      = "FF2563EB"   # 파란색 (S등급)
CLR_A      = "FF16A34A"   # 초록색 (A등급)
CLR_B      = "FFD97706"   # 주황색 (B등급)
CLR_C      = "FF6B7280"   # 회색   (C등급)
CLR_YELLOW = "FFFFF8E7"   # 연한 노란 배경 (공고명·공고번호)
CLR_HEADER = "FF1E293B"   # 헤더 배경 (다크)
CLR_GUIDE  = "FFF0F9FF"   # 사용법 시트 배경

GRADE_COLORS = {
    "S": ("FFDBEAFE", "FF1D4ED8"),   # (배경, 글자)
    "A": ("FFD1FAE5", "FF065F46"),
    "B": ("FFFEF3C7", "FF92400E"),
    "C": ("FFF3F4F6", "FF374151"),
}

_THIN = Side(style="thin", color="FFD1D5DB")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


# ─────────────────────────────────────────────
# 헬퍼
# ─────────────────────────────────────────────

def _fill(hex_color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=hex_color)


def _font(bold=False, size=10, color="FF111827", name="맑은 고딕") -> Font:
    return Font(bold=bold, size=size, color=color, name=name)


def _center(wrap=False) -> Alignment:
    return Alignment(horizontal="center", vertical="center", wrap_text=wrap)


def _left(wrap=False) -> Alignment:
    return Alignment(horizontal="left", vertical="center", wrap_text=wrap)


def _fmt_budget(raw: str) -> str:
    """추정가격 → '억' 단위 문자열. 예: '27272727' → '0.27억'"""
    try:
        val = float(str(raw).replace(",", "").replace(" ", "") or "0")
        if val == 0:
            return "미공개"
        return f"{val / 1e8:.2f}억"
    except (ValueError, TypeError):
        return "미공개"


def _fmt_date(raw: str) -> str:
    dt = parse_date(raw)
    return dt.strftime("%Y-%m-%d") if dt else "-"


def _days_left(raw: str, today: datetime) -> str:
    dt = parse_date(raw)
    if not dt:
        return "-"
    d = (dt - today).days
    if d < 0:
        return "마감"
    if d == 0:
        return "D-Day"
    return f"D-{d}"


def _matched_code_label(codes: list[str]) -> str:
    parts = []
    for c in codes:
        name = CODE_NAMES.get(c, c)
        parts.append(f"{c} {name}")
    return "\n".join(parts) if parts else "-"


def _g2b_deep_link(bid: dict) -> str:
    bid_no = bid.get("bidNtceNo", "")
    bid_ord = bid.get("bidNtceOrd", "000") or "000"
    return (
        "https://www.g2b.go.kr/link/PNPE027_01/single/"
        f"?bidPbancNo={bid_no}&bidPbancOrd={bid_ord}"
    )


# ─────────────────────────────────────────────
# 시트 빌더 — 키워드별 공고 시트
# ─────────────────────────────────────────────

_COL_HEADERS = [
    ("순위",        5),
    ("등급",        6),
    ("점수",        6),
    ("공고명 (복사용)", 50),
    ("발주기관",    20),
    ("입찰방식",    16),
    ("예산",        10),
    ("마감일",      12),
    ("D-Day",        8),
    ("매칭 등록업종", 28),
    ("자격출처",    14),
    ("적합 이유",   30),
    ("공고번호 (복사용)", 22),
    ("공고상세",    12),
]


def _build_keyword_sheet(ws, keyword: str, bids: list[dict],
                          today: datetime, top_n: int = 10):
    ws.title = f"🔍 {keyword}"

    # 행 1: 키워드 + 건수 제목
    ws.merge_cells("A1:N1")
    c = ws["A1"]
    c.value = f"🔍  {keyword}  —  자격 통과 공고 상위 {min(top_n, len(bids))}건 / 전체 {len(bids)}건"
    c.font      = _font(bold=True, size=13, color="FFFFFFFF")
    c.fill      = _fill(CLR_HEADER)
    c.alignment = _center()
    ws.row_dimensions[1].height = 28

    # 행 2: 나라장터 링크
    ws.merge_cells("A2:N2")
    c = ws["A2"]
    c.value     = "🔗  공고 행의 '공고상세' 링크를 클릭하면 나라장터 상세 화면으로 이동합니다."
    c.font      = Font(bold=True, size=10, color="FF2563EB",
                       underline="single", name="맑은 고딕")
    c.fill      = _fill("FFE0F2FE")
    c.alignment = _center()
    ws.row_dimensions[2].height = 22
    ws.cell(row=2, column=1).hyperlink = "https://www.g2b.go.kr/"

    # 행 3: 컬럼 헤더
    for col_i, (header, width) in enumerate(_COL_HEADERS, start=1):
        c = ws.cell(row=3, column=col_i)
        c.value     = header
        c.font      = _font(bold=True, size=9, color="FFFFFFFF")
        c.fill      = _fill("FF334155")
        c.alignment = _center(wrap=True)
        c.border    = _BORDER
        ws.column_dimensions[get_column_letter(col_i)].width = width
    ws.row_dimensions[3].height = 30

    # 공고 데이터 행 (상위 top_n)
    for rank, bid in enumerate(bids[:top_n], start=1):
        row = rank + 3
        grade = bid.get("_grade", "C")
        g_bg, g_fg = GRADE_COLORS.get(grade, ("FFF3F4F6", "FF374151"))

        def wc(col_i, value, bold=False, bg=None, fg="FF111827",
               align=None, font_name="맑은 고딕", wrap=False):
            c = ws.cell(row=row, column=col_i)
            c.value     = value
            c.font      = Font(bold=bold, size=9, color=fg, name=font_name)
            c.fill      = _fill(bg) if bg else PatternFill()
            c.alignment = align or _left(wrap=wrap)
            c.border    = _BORDER
            return c

        wc(1,  rank,                               align=_center())
        # 등급 셀: 배경색 적용
        gc = ws.cell(row=row, column=2)
        gc.value     = grade
        gc.font      = Font(bold=True, size=10, color=g_fg, name="맑은 고딕")
        gc.fill      = _fill(g_bg)
        gc.alignment = _center()
        gc.border    = _BORDER

        wc(3,  bid.get("_score", 0),               align=_center())
        wc(4,  bid.get("bidNtceNm", ""),            bold=True,
                                                    bg=CLR_YELLOW,
                                                    wrap=True)
        wc(5,  bid.get("ntceInsttNm", ""),          wrap=True)
        wc(6,  bid.get("bidMthdNm", ""),            wrap=True)
        wc(7,  _fmt_budget(bid.get("presmptPrce",
                           bid.get("asignBdgtAmt", ""))),
                                                     align=_center())
        wc(8,  _fmt_date(bid.get("bidClseDt", "")),align=_center())
        wc(9,  _days_left(bid.get("bidClseDt", ""), today),
                                                     align=_center())
        wc(10, _matched_code_label(bid.get("_matched_codes", [])),
                                                    wrap=True)
        wc(11, bid.get("_license_source_label", ""),align=_center())
        wc(12, " / ".join(bid.get("_reasons", [])), wrap=True)
        # 공고번호 셀: 노란 배경 + Consolas
        nc = ws.cell(row=row, column=13)
        nc.value     = bid.get("bidNtceFullNo", bid.get("bidNtceNo", ""))
        nc.font      = Font(bold=True, size=9, color="FF1E293B",
                            name="Consolas")
        nc.fill      = _fill(CLR_YELLOW)
        nc.alignment = _center()
        nc.border    = _BORDER

        lc = ws.cell(row=row, column=14)
        lc.value     = "열기"
        lc.hyperlink = _g2b_deep_link(bid)
        lc.font      = Font(bold=True, size=9, color="FF2563EB",
                            underline="single", name="맑은 고딕")
        lc.alignment = _center()
        lc.border    = _BORDER

        ws.row_dimensions[row].height = 40

    # 틀 고정 (헤더 3행 + 컬럼1 고정)
    ws.freeze_panes = "B4"


# ─────────────────────────────────────────────
# 시트 빌더 — 자격 미해당 시트
# ─────────────────────────────────────────────

_DQ_HEADERS = [
    ("공고명",     45),
    ("발주기관",   18),
    ("입찰방식",   14),
    ("마감일",     12),
    ("미해당 사유", 35),
    ("공고번호",   22),
    ("공고상세",   12),
]


def _build_disqualified_sheet(ws, disqualified: list[dict]):
    ws.title = "🚫 자격 미해당"

    ws.merge_cells("A1:G1")
    c = ws["A1"]
    c.value     = f"🚫  자격 미해당 공고  —  총 {len(disqualified)}건"
    c.font      = _font(bold=True, size=12, color="FFFFFFFF")
    c.fill      = _fill("FF991B1B")
    c.alignment = _center()
    ws.row_dimensions[1].height = 26

    for col_i, (header, width) in enumerate(_DQ_HEADERS, start=1):
        c = ws.cell(row=2, column=col_i)
        c.value     = header
        c.font      = _font(bold=True, size=9, color="FFFFFFFF")
        c.fill      = _fill("FF334155")
        c.alignment = _center()
        c.border    = _BORDER
        ws.column_dimensions[get_column_letter(col_i)].width = width
    ws.row_dimensions[2].height = 24

    for row_i, bid in enumerate(disqualified, start=3):
        def wc(col_i, value, wrap=False):
            c = ws.cell(row=row_i, column=col_i)
            c.value     = value
            c.font      = _font(size=9)
            c.alignment = _left(wrap=wrap)
            c.border    = _BORDER

        wc(1, bid.get("bidNtceNm", ""),                           wrap=True)
        wc(2, bid.get("ntceInsttNm", ""))
        wc(3, bid.get("bidMthdNm", ""))
        wc(4, _fmt_date(bid.get("bidClseDt", "")))
        wc(5, " / ".join(bid.get("_disqualify_reasons", [])),     wrap=True)

        nc = ws.cell(row=row_i, column=6)
        nc.value     = bid.get("bidNtceFullNo", bid.get("bidNtceNo", ""))
        nc.font      = Font(size=9, name="Consolas")
        nc.alignment = _center()
        nc.border    = _BORDER

        lc = ws.cell(row=row_i, column=7)
        lc.value     = "열기"
        lc.hyperlink = _g2b_deep_link(bid)
        lc.font      = Font(bold=True, size=9, color="FF2563EB",
                            underline="single", name="맑은 고딕")
        lc.alignment = _center()
        lc.border    = _BORDER

        ws.row_dimensions[row_i].height = 35

    ws.freeze_panes = "A3"


# ─────────────────────────────────────────────
# 시트 빌더 — 사용법 안내 시트
# ─────────────────────────────────────────────

def _build_guide_sheet(ws, keyword_stats: dict[str, tuple[int, int]], scan_dt: datetime):
    """
    keyword_stats: {keyword: (qualified_count, total_count)}
    """
    ws.title = "📖 사용법 안내"
    ws.sheet_view.showGridLines = False

    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 50
    ws.column_dimensions["D"].width = 18

    def mw(row, col, end_col, value, bold=False, size=10,
           fg="FF111827", bg=None, align=None, wrap=False):
        """merge + write"""
        if end_col > col:
            ws.merge_cells(
                start_row=row, start_column=col,
                end_row=row,   end_column=end_col
            )
        c = ws.cell(row=row, column=col)
        c.value     = value
        c.font      = Font(bold=bold, size=size, color=fg, name="맑은 고딕")
        if bg:
            c.fill  = _fill(bg)
        c.alignment = align or _left(wrap=wrap)
        ws.row_dimensions[row].height = 22
        return c

    r = 1
    mw(r, 1, 4, f"📊  핀인사이트 나라장터 공고 스캐너  —  스캔일시: {scan_dt.strftime('%Y-%m-%d %H:%M')}",
       bold=True, size=14, fg="FFFFFFFF", bg=CLR_HEADER, align=_center())
    ws.row_dimensions[r].height = 34

    r += 1
    mw(r, 1, 4, "", bg=CLR_GUIDE)

    r += 1
    mw(r, 1, 4, "🔎  나라장터에서 공고 찾는 방법", bold=True, size=12, fg="FF1D4ED8", bg=CLR_GUIDE)
    ws.row_dimensions[r].height = 26

    steps = [
        ("STEP 1", "공고상세 링크 클릭",
         "각 키워드 시트의 공고상세 열에서 '열기' 링크를 클릭하세요."),
        ("STEP 2", "공고번호 또는 공고명 복사",
         "필요하면 노란 배경 셀(공고번호·공고명)을 복사하세요."),
        ("STEP 3", "나라장터에서 상세 확인",
         "링크가 열리지 않을 때는 공고번호를 나라장터 검색창에 붙여넣어 확인하세요."),
    ]
    for step, title, desc in steps:
        r += 1
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=2)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        ws.cell(row=r, column=2).value     = f"  {step}  {title}"
        ws.cell(row=r, column=2).font      = Font(bold=True, size=10, color="FF1D4ED8", name="맑은 고딕")
        ws.cell(row=r, column=2).fill      = _fill("FFE0F2FE")
        ws.cell(row=r, column=2).alignment = _left()
        ws.cell(row=r, column=3).value     = desc
        ws.cell(row=r, column=3).font      = Font(size=9, name="맑은 고딕")
        ws.cell(row=r, column=3).fill      = _fill(CLR_GUIDE)
        ws.cell(row=r, column=3).alignment = _left()
        ws.row_dimensions[r].height = 22

    r += 1
    link_cell = ws.cell(row=r, column=2)
    link_cell.value     = "  🔗  나라장터 바로 열기"
    link_cell.hyperlink = "https://www.g2b.go.kr/"
    link_cell.font      = Font(bold=True, size=10, color="FF2563EB",
                               underline="single", name="맑은 고딕")
    link_cell.fill      = _fill(CLR_GUIDE)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    ws.row_dimensions[r].height = 24

    r += 2
    mw(r, 1, 4, "💡  검색 팁", bold=True, size=12, fg="FF1D4ED8", bg=CLR_GUIDE)
    ws.row_dimensions[r].height = 26

    tips = [
        "공고상세 링크는 bidPbancNo와 bidPbancOrd를 사용해 생성됩니다.",
        "공고번호(예: R26BK01512858)가 가장 정확한 검색 키입니다.",
        "수동 검색 시 차수(-000)는 빼고 검색하세요 (재공고 시 차수가 바뀝니다).",
        "공고명이 길면 핵심 키워드 일부만 검색해도 됩니다.",
        "자격출처 '검색조건(API)' = 업종코드 검색조건 매칭 / '추정(사업내용)' = 공고명 분석.",
    ]
    for tip in tips:
        r += 1
        mw(r, 2, 4, f"  ▸  {tip}", size=9, bg=CLR_GUIDE)

    r += 2
    mw(r, 1, 4, "🏢  핀인사이트 등록업종 (자격 필터 기준)", bold=True, size=12, fg="FF1D4ED8", bg=CLR_GUIDE)
    ws.row_dimensions[r].height = 26

    codes = [
        ("1468", "SW(컴퓨터관련서비스)", "시스템 구축·SI·SW 개발·플랫폼 구축·운영"),
        ("1469", "SW(디지털콘텐츠)",     "AI·챗봇·디지털 콘텐츠 개발·교육 콘텐츠"),
        ("1470", "SW(DB제작/검색)",      "DB 구축·자료처리·데이터 분석·검색시스템"),
        ("1169", "학술연구용역",          "정책연구·R&D·타당성 검토·로드맵·교육평가"),
        ("9999", "기타자유업종",          "면허 제한 없는 일반 용역"),
    ]
    for code, name, desc in codes:
        r += 1
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=2)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=3)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=4)
        ws.cell(row=r, column=2).value     = f"  {code}"
        ws.cell(row=r, column=2).font      = Font(bold=True, size=9, color="FF1D4ED8", name="Consolas")
        ws.cell(row=r, column=2).fill      = _fill("FFE0F2FE")
        ws.cell(row=r, column=3).value     = name
        ws.cell(row=r, column=3).font      = Font(bold=True, size=9, name="맑은 고딕")
        ws.cell(row=r, column=3).fill      = _fill(CLR_GUIDE)
        ws.cell(row=r, column=4).value     = desc
        ws.cell(row=r, column=4).font      = Font(size=9, name="맑은 고딕")
        ws.cell(row=r, column=4).fill      = _fill(CLR_GUIDE)
        ws.row_dimensions[r].height = 20

    r += 2
    mw(r, 1, 4, "📈  키워드별 스캔 결과 요약", bold=True, size=12, fg="FF1D4ED8", bg=CLR_GUIDE)
    ws.row_dimensions[r].height = 26

    r += 1
    for col_i, header in enumerate(["키워드", "자격 통과", "전체 수집", "통과율"], start=2):
        c = ws.cell(row=r, column=col_i)
        c.value     = header
        c.font      = _font(bold=True, size=9, color="FFFFFFFF")
        c.fill      = _fill("FF334155")
        c.alignment = _center()
        c.border    = _BORDER
    ws.row_dimensions[r].height = 22

    for kw, (q_cnt, t_cnt) in keyword_stats.items():
        r += 1
        rate = f"{q_cnt/t_cnt*100:.0f}%" if t_cnt else "-"
        for col_i, value in enumerate([f"🔍 {kw}", q_cnt, t_cnt, rate], start=2):
            c = ws.cell(row=r, column=col_i)
            c.value     = value
            c.font      = _font(size=9)
            c.fill      = _fill(CLR_GUIDE)
            c.alignment = _center()
            c.border    = _BORDER
        ws.row_dimensions[r].height = 20


# ─────────────────────────────────────────────
# 메인 리포트 생성 함수
# ─────────────────────────────────────────────

def generate_report(
    keyword_results: dict[str, dict],
    output_path: str,
    today: datetime = None,
    top_n: int = 10,
) -> str:
    """
    xlsx 리포트 생성.

    Args:
        keyword_results: {
            keyword: {
                'qualified':    [bid, ...],   # 자격 통과 + 점수화 완료
                'disqualified': [bid, ...],   # 자격 미해당
                'total':        int,          # 전체 수집 건수
            }
        }
        output_path: 저장 경로 (.xlsx)
        today: 기준일 (None이면 현재)
        top_n: 키워드 시트당 표시 건수

    Returns:
        저장된 파일 경로
    """
    if today is None:
        today = datetime.now()

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)   # 기본 시트 제거

    # 1. 사용법 안내 시트 (첫 번째)
    ws_guide = wb.create_sheet()
    keyword_stats = {
        kw: (len(v["qualified"]), v["total"])
        for kw, v in keyword_results.items()
    }
    _build_guide_sheet(ws_guide, keyword_stats, today)

    # 2. 키워드별 공고 시트
    for keyword, v in keyword_results.items():
        ws = wb.create_sheet()
        _build_keyword_sheet(ws, keyword, v["qualified"], today, top_n=top_n)

    # 3. 자격 미해당 시트 (마지막)
    all_disqualified = []
    for v in keyword_results.values():
        all_disqualified.extend(v["disqualified"])
    # 중복 제거
    seen_nos = set()
    unique_dq = []
    for bid in all_disqualified:
        no = bid.get("bidNtceNo", "")
        if no not in seen_nos:
            seen_nos.add(no)
            unique_dq.append(bid)

    ws_dq = wb.create_sheet()
    _build_disqualified_sheet(ws_dq, unique_dq)

    wb.save(output_path)
    return output_path
